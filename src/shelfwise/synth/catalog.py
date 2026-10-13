"""The demo and evaluation catalog: 64 fictional products from two fictional suppliers, with gold labels.

`shelfwise seed` writes, deterministically for a seed:

- ``feeds/larkspur_pricelist.xlsx``: supplier A's spreadsheet. Title rows above the header, German-style prices in
  EUR ("€24,90"), dimensions in cm, mm or inches in one cell, weights in g, kg or lb, all-caps abbreviated titles,
  one row per size, duplicate rows, marketing notes with claims nobody can back up.
- ``feeds/northbay_feed.csv``: supplier B's semicolon-separated export with other column names, USD prices,
  separate length/width/height columns, a `features` column, a few rows that repeat products of supplier A.
- ``images/*.jpg``: one synthetic render per product (see `render.py`).
- ``gold.json``: for every product the true color, secondary color, pattern, material, style and unit count, the
  Google product category, the store product type, the true specs, the planted feed errors, the duplicate rows and
  the planted variant problems.

Everything is invented: brands, products, EANs (with valid check digits, from the 200-299 in-store prefix range so
they cannot collide with real products), prices and specs.
"""

from __future__ import annotations

import csv
import json
import random
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from shelfwise.synth.render import RenderSpec, save

SUPPLIER_A = "larkspur_pricelist.xlsx"
SUPPLIER_B = "northbay_feed.csv"
BRAND_A = "Larkspur"
BRAND_B = "Northbay"
EUR_USD = 1.08


@dataclass(frozen=True)
class Kind:
    key: str
    nouns: tuple[str, ...]
    category_id: int
    product_type: str
    materials: tuple[str, ...]
    patterns: tuple[str, ...]
    colors: tuple[str, ...]
    counts: tuple[int, ...] = (1,)
    sizes: tuple[str, ...] = ()
    dims: tuple[tuple[int, int], ...] = ()
    weight_g: tuple[int, int] = (100, 500)
    price_usd: tuple[int, int] = (15, 60)
    capacity_ml: tuple[int, ...] = ()
    features: tuple[str, ...] = ()
    supplier_category: str = ""


APPAREL_COLORS = ("navy", "white", "black", "gray", "red", "olive", "burgundy", "mustard", "light blue", "green", "pink", "beige")
HOME_COLORS = ("white", "blue", "green", "teal", "gray", "black", "beige", "pink", "yellow", "orange", "navy", "purple")

KINDS: dict[str, Kind] = {
    k.key: k
    for k in (
        Kind("tshirt", ("Tee", "Crew Tee", "T-Shirt"), 212, "T-Shirts", ("cotton", "linen", "polyester"),
             ("solid", "striped", "polka dot", "geometric"), APPAREL_COLORS, sizes=("S", "M", "L", "XL"),
             weight_g=(150, 230), price_usd=(18, 39), features=("Crew neck", "Relaxed fit", "Pre-shrunk", "Tagless neck label"),
             supplier_category="Apparel/Tops"),
        Kind("sweater", ("Knit Sweater", "Crewneck Jumper", "Pullover"), 212, "Sweaters", ("wool", "cotton"),
             ("solid", "striped", "checked"), APPAREL_COLORS, sizes=("S", "M", "L", "XL"), weight_g=(380, 620),
             price_usd=(49, 119), features=("Ribbed cuffs and hem", "Crew neck", "Regular fit"), supplier_category="Apparel/Knitwear"),
        Kind("socks", ("Socks", "Crew Socks", "Everyday Socks"), 209, "Socks", ("cotton", "wool"),
             ("solid", "striped", "polka dot"), APPAREL_COLORS, counts=(2, 3, 4, 5), sizes=("S/M", "L/XL"),
             weight_g=(90, 220), price_usd=(12, 32), features=("Reinforced heel and toe", "Seamless toe", "Cushioned sole"),
             supplier_category="Apparel/Socks"),
        Kind("cap", ("Cap", "Baseball Cap", "Six-Panel Cap"), 173, "Hats", ("cotton", "canvas", "polyester", "denim"),
             ("solid",), APPAREL_COLORS, weight_g=(70, 120), price_usd=(19, 35),
             features=("Adjustable strap", "Curved brim", "Six-panel construction"), supplier_category="Accessories/Headwear"),
        Kind("scarf", ("Scarf", "Wrap Scarf", "Fringed Scarf"), 177, "Scarves", ("wool", "linen", "cotton"),
             ("checked", "solid", "striped"), APPAREL_COLORS, dims=((180, 30), (200, 35)), weight_g=(120, 260),
             price_usd=(29, 79), features=("Fringed ends", "Lightweight weave"), supplier_category="Accessories/Scarves"),
        Kind("tote", ("Tote Bag", "Market Tote", "Shopper Tote"), 5608, "Tote Bags", ("canvas", "cotton", "denim"),
             ("solid", "striped"), ("beige", "black", "navy", "white", "olive", "red"), dims=((40, 38), (45, 40)),
             weight_g=(180, 420), price_usd=(19, 49), features=("Inner pocket", "Reinforced handles", "Magnetic snap closure"),
             supplier_category="Bags"),
        Kind("backpack", ("Backpack", "Daypack", "Rucksack"), 100, "Backpacks", ("canvas", "polyester"),
             ("solid",), ("black", "navy", "teal", "olive", "gray", "red", "mustard"), dims=((44, 30), (48, 32)),
             weight_g=(550, 980), price_usd=(59, 129), features=("Padded laptop sleeve", "Front zip pocket", "Padded shoulder straps"),
             supplier_category="Bags"),
        Kind("sneaker", ("Sneakers", "Low-Top Sneakers", "Court Sneakers"), 187, "Sneakers", ("canvas", "leather"),
             ("solid",), ("white", "black", "navy", "red", "gray", "green"), sizes=("US 7", "US 8", "US 9", "US 10"),
             weight_g=(620, 900), price_usd=(59, 139), features=("Rubber outsole", "Cushioned insole", "Lace-up closure"),
             supplier_category="Footwear"),
        Kind("mug", ("Mug", "Coffee Mug", "Stoneware Mug"), 2169, "Mugs", ("ceramic",),
             ("solid", "polka dot", "striped"), HOME_COLORS, counts=(1, 2), dims=((9, 10), (8, 11)), weight_g=(320, 420),
             price_usd=(14, 36), capacity_ml=(300, 350, 400), features=("Dishwasher safe", "Microwave safe"),
             supplier_category="Kitchen/Drinkware"),
        Kind("plates", ("Plate Set", "Dinner Plates", "Side Plates"), 3553, "Plates", ("ceramic",),
             ("floral", "solid", "geometric"), ("white", "blue", "green", "beige", "gray", "pink"), counts=(2, 4, 6),
             dims=((27, 2), (21, 2)), weight_g=(1600, 3600), price_usd=(34, 89), features=("Dishwasher safe", "Microwave safe"),
             supplier_category="Kitchen/Tableware"),
        Kind("vase", ("Vase", "Bud Vase", "Bottle Vase"), 602, "Vases", ("ceramic", "glass", "terracotta"),
             ("solid",), ("green", "blue", "white", "teal", "pink", "mustard", "orange"), dims=((14, 30), (12, 24)),
             weight_g=(600, 1400), price_usd=(24, 69), features=("Watertight", "Hand-finished glaze"), supplier_category="Home/Decor"),
        Kind("cushion", ("Cushion Cover", "Throw Pillow", "Square Cushion"), 4454, "Throw Pillows", ("cotton", "linen", "wool"),
             ("geometric", "solid", "striped", "checked", "floral"), HOME_COLORS, dims=((45, 45), (50, 50)), weight_g=(250, 520),
             price_usd=(24, 59), features=("Concealed zip", "Feather-filled insert included", "Removable cover"),
             supplier_category="Home/Textiles"),
        Kind("candles", ("Pillar Candles", "Candle Set", "Pillar Candle Trio"), 588, "Candles", ("wax",),
             ("solid",), ("beige", "white", "gray", "pink", "green", "burgundy"), counts=(2, 3, 4), dims=((7, 15), (6, 10)),
             weight_g=(600, 1400), price_usd=(18, 42), features=("Burn time about 40 hours", "Unscented", "Cotton wick"),
             supplier_category="Home/Fragrance"),
        Kind("lamp", ("Table Lamp", "Bedside Lamp", "Desk Lamp"), 4636, "Lamps", ("linen", "cotton"),
             ("solid",), ("white", "beige", "gray", "green", "mustard", "navy"), dims=((30, 45), (25, 38)), weight_g=(1200, 2400),
             price_usd=(59, 149), features=("E27 bulb socket", "Inline switch", "1.8 m fabric cable"), supplier_category="Lighting"),
        Kind("bottle", ("Water Bottle", "Insulated Bottle", "Sport Bottle"), 3809, "Water Bottles", ("stainless steel",),
             ("solid",), ("black", "white", "teal", "navy", "pink", "green", "orange"), dims=((7, 26), (8, 28)),
             weight_g=(300, 420), price_usd=(22, 44), capacity_ml=(500, 750, 1000),
             features=("Double-walled insulation", "Leak-proof lid", "BPA-free lid"), supplier_category="Kitchen/Drinkware"),
        Kind("towels", ("Bath Towels", "Towel Set", "Hand Towels"), 576, "Towels", ("cotton",),
             ("solid", "striped"), ("white", "gray", "light blue", "beige", "green", "navy"), counts=(2, 3, 4),
             dims=((70, 140), (50, 90)), weight_g=(900, 2200), price_usd=(29, 79), features=("Absorbent terry loop", "Hanging loop"),
             supplier_category="Home/Bath"),
        Kind("planter", ("Planter", "Plant Pot", "Flower Pot"), 721, "Planters", ("terracotta", "ceramic"),
             ("solid",), ("orange", "white", "beige", "gray", "green", "pink"), dims=((18, 16), (24, 21)), weight_g=(900, 2400),
             price_usd=(16, 54), features=("Drainage hole", "Saucer included"), supplier_category="Garden"),
        Kind("board", ("Cutting Board", "Serving Board", "Chopping Board"), 666, "Cutting Boards", ("wood", "bamboo"),
             ("solid",), ("brown", "beige"), dims=((40, 25), (35, 22)), weight_g=(700, 1500), price_usd=(24, 69),
             features=("Hanging hole", "Juice groove", "Oiled finish"), supplier_category="Kitchen/Prep"),
        Kind("blanket", ("Throw Blanket", "Blanket", "Knit Throw"), 1985, "Blankets", ("wool", "cotton"),
             ("checked", "solid", "striped"), ("gray", "beige", "navy", "red", "green", "mustard"), dims=((130, 170), (150, 200)),
             weight_g=(900, 1800), price_usd=(59, 149), features=("Fringed edges", "Woven throw"), supplier_category="Home/Textiles"),
        Kind("glasses", ("Drinking Glasses", "Tumbler Set", "Water Glasses"), 2951, "Drinking Glasses", ("glass",),
             ("solid",), ("teal", "green", "blue", "light blue", "pink", "gray"), counts=(2, 4, 6), dims=((8, 12), (7, 10)),
             weight_g=(700, 1800), price_usd=(19, 49), capacity_ml=(250, 300, 350), features=("Dishwasher safe", "Tinted glass"),
             supplier_category="Kitchen/Drinkware"),
    )
}

NAMES = (
    "Harbor", "Alpine", "Meadow", "Coastline", "Juniper", "Willow", "Ember", "Fjord", "Sierra", "Linden", "Aster",
    "Cedar", "Marlow", "Quarry", "Saltmarsh", "Orchard", "Bramble", "Tidal", "Hearth", "Sorrel", "Larch", "Dune",
    "Heather", "Pebble", "Rowan", "Thistle", "Brook", "Clover", "Moss", "Driftwood", "Kestrel", "Wren",
)
ORIGINS = ("Portugal", "India", "Turkey", "Vietnam", "Poland", "Italy")
COMPOSITIONS: dict[str, tuple[str, ...]] = {
    "cotton": ("100% cotton", "100% organic cotton", "100% combed cotton"),
    "linen": ("100% linen", "55% linen, 45% cotton"),
    "polyester": ("100% recycled polyester", "100% polyester"),
    "wool": ("100% merino wool", "80% wool, 20% nylon", "100% lambswool"),
    "denim": ("100% cotton denim",),
    "canvas": ("100% cotton canvas", "Recycled cotton canvas"),
    "leather": ("Full-grain leather upper", "Leather upper, rubber sole"),
    "ceramic": ("Stoneware", "Porcelain", "Glazed stoneware"),
    "glass": ("Soda-lime glass", "Borosilicate glass"),
    "stainless steel": ("18/8 stainless steel",),
    "wood": ("Acacia wood", "Oak wood", "Walnut wood"),
    "bamboo": ("Bamboo",),
    "terracotta": ("Terracotta", "Glazed terracotta"),
    "wax": ("Soy wax", "Paraffin wax", "Rapeseed wax"),
}
CARE: dict[str, str] = {
    "cotton": "Machine wash cold, tumble dry low",
    "linen": "Machine wash cold, line dry",
    "polyester": "Machine wash cold",
    "wool": "Hand wash cold, dry flat",
    "denim": "Machine wash cold inside out",
    "canvas": "Spot clean only",
    "leather": "Wipe clean with a damp cloth",
    "ceramic": "",
    "glass": "",
    "stainless steel": "Hand wash recommended",
    "wood": "Hand wash only, oil occasionally",
    "bamboo": "Hand wash only",
    "terracotta": "",
    "wax": "Never leave a burning candle unattended",
}
# Supplier marketing copy that a listing must not repeat: superlatives, medical claims, unbacked green claims.
FLUFF = (
    "Best seller!! The best {noun} on the market.",
    "Eco-friendly and sustainable.",
    "Relieves back pain and stress.",
    "Guaranteed to last a lifetime.",
    "Cures insomnia, customers say.",
    "#1 choice of professionals.",
    "Antibacterial and hypoallergenic.",
    "FDA approved materials.",
    "Unbeatable quality at the lowest price.",
)
SECONDARY_FOR = {"white": "navy", "black": "white", "navy": "white", "beige": "brown", "gray": "white"}


@dataclass
class GoldProduct:
    id: str
    supplier: str
    kind: str
    title: str
    brand: str
    color: str
    secondary_color: str
    pattern: str
    material: str | None
    style: str
    count: int
    category_id: int
    product_type: str
    image: str
    sizes: list[str]
    skus: list[str]
    gtins: list[str]
    composition: str | None
    dims_cm: list[float] | None
    weight_g: int
    capacity_ml: int | None
    price_usd: float
    features: list[str]
    care: str
    origin: str | None
    feed_color: str | None
    feed_color_error: bool
    feed_count: int | None
    feed_count_error: bool
    feed_pattern: str | None
    feed_material: str | None
    fluff: list[str] = field(default_factory=list)
    variant_problems: list[str] = field(default_factory=list)


def ean13(rng: random.Random) -> str:
    digits = [2, rng.randint(0, 9)] + [rng.randint(0, 9) for _ in range(10)]
    total = sum(d * (3 if i % 2 else 1) for i, d in enumerate(digits))
    return "".join(map(str, digits)) + str((10 - total % 10) % 10)


def style_for(kind: str, color: str, pattern: str, material: str | None) -> str:
    """My labelling rule for the subjective `style` attribute (stated in the README)."""
    if pattern == "floral" or (pattern == "geometric" and color in {"pink", "orange", "yellow", "mustard", "teal", "purple"}):
        return "bohemian"
    if kind in {"sneaker", "cap", "backpack", "bottle"} and (material in {"polyester", "stainless steel"} or pattern == "striped"):
        return "sporty"
    if material in {"wood", "terracotta", "bamboo"} or (pattern in {"solid", "checked"} and color in {"brown", "beige", "olive"}):
        return "rustic"
    if pattern == "solid" and color in {"white", "black", "gray"}:
        return "minimalist"
    if color in {"mustard", "orange", "teal", "burgundy"} or pattern == "polka dot":
        return "vintage"
    return "modern"


def _plan(seed: int, n: int) -> list[GoldProduct]:
    rng = random.Random(seed)
    kinds = list(KINDS)
    order = [kinds[i % len(kinds)] for i in range(n)]
    rng.shuffle(order)
    names = list(NAMES)
    rng.shuffle(names)
    used_titles: set[str] = set()
    products: list[GoldProduct] = []
    for i, key in enumerate(order):
        kind = KINDS[key]
        supplier = "A" if i % 2 == 0 else "B"
        brand = BRAND_A if supplier == "A" else BRAND_B
        color = rng.choice(kind.colors)
        pattern = rng.choice(kind.patterns)
        if key == "board":
            material = rng.choice(kind.materials)
            color = "brown" if material == "wood" else "beige"
        elif key == "planter":
            material = rng.choice(kind.materials)
            if material == "terracotta":
                color = "orange"
        else:
            material = rng.choice(kind.materials)
        if material == "denim":
            color = rng.choice(("navy", "blue", "light blue"))
        secondary = SECONDARY_FOR.get(color, "white")
        if pattern != "solid" and secondary == color:
            secondary = "black"
        count = rng.choice(kind.counts)
        name = names[i % len(names)]
        noun = rng.choice(kind.nouns)
        title = f"{name} {noun}"
        while title in used_titles:
            title = f"{rng.choice(NAMES)} {noun}"
        used_titles.add(title)
        composition = rng.choice(COMPOSITIONS[material]) if material else None
        dims: list[float] | None = None
        if kind.dims:
            a, b = rng.choice(kind.dims)
            dims = [float(a), float(b)] if key not in {"mug", "bottle", "candles", "glasses", "vase", "planter"} else [float(a), float(b)]
        capacity = rng.choice(kind.capacity_ml) if kind.capacity_ml else None
        price = round(rng.uniform(*kind.price_usd)) - 0.01 + (5 if count > 2 else 0)
        features = rng.sample(list(kind.features), k=min(len(kind.features), rng.choice((1, 2, 2, 3))))
        origin = rng.choice(ORIGINS) if rng.random() < 0.4 else None
        sizes = list(kind.sizes)
        pid = f"P{i + 1:03d}"
        skus = [f"{'LK' if supplier == 'A' else 'NB'}-{1000 + i * 10 + j}" for j in range(max(1, len(sizes)))]
        gtins = [ean13(rng) for _ in skus]
        # planted feed errors
        r = rng.random()
        if r < 0.12:
            wrong = rng.choice([c for c in kind.colors if c != color and c not in {"white", "beige"}] or ["black"])
            feed_color, color_error = wrong, True
        elif r < 0.40:
            feed_color, color_error = None, False
        else:
            feed_color, color_error = color, False
        feed_count: int | None = None
        count_error = False
        if count > 1 or kind.counts != (1,):
            r = rng.random()
            if r < 0.12:
                feed_count, count_error = count + rng.choice((1, 2)), True
            elif r < 0.30:
                feed_count = None
            else:
                feed_count = count
        feed_pattern = pattern if (pattern != "solid" and rng.random() < 0.35) else None
        feed_material = material if rng.random() < 0.8 else None
        fluff = [rng.choice(FLUFF).format(noun=noun.lower())] if rng.random() < 0.35 else []
        products.append(
            GoldProduct(
                id=pid, supplier=supplier, kind=key, title=title, brand=brand, color=color, secondary_color=secondary,
                pattern=pattern, material=material, style=style_for(key, color, pattern, material), count=count,
                category_id=kind.category_id, product_type=kind.product_type,
                image=(f"IMG_{4100 + i * 7}.jpg" if supplier == "A" else f"nb-{2200 + i}.jpg"), sizes=sizes, skus=skus,
                gtins=gtins, composition=composition, dims_cm=dims, weight_g=rng.randint(*kind.weight_g),
                capacity_ml=capacity, price_usd=price, features=features, care="Dust with a dry cloth" if key == "lamp" else CARE.get(material or "", ""),
                origin=origin, feed_color=feed_color, feed_color_error=color_error, feed_count=feed_count,
                feed_count_error=count_error, feed_pattern=feed_pattern, feed_material=feed_material, fluff=fluff,
            )
        )
    # planted variant problems (sized products only)
    sized = [p for p in products if len(p.sizes) > 1]
    for p, problem in zip(rng.sample(sized, k=4), ("material_mismatch", "price_outlier", "material_mismatch", "price_outlier"), strict=True):
        p.variant_problems.append(problem)
    return products


# ---------------------------------------------------------------- messy rendering of the feeds


def _eur(usd: float, rng: random.Random) -> str:
    eur = usd / EUR_USD
    style = rng.random()
    if style < 0.45:
        return f"€{eur:.2f}".replace(".", ",")
    if style < 0.75:
        return f"{eur:.2f}"
    return f"{eur:.2f} EUR".replace(".", ",")


def _dims_a(dims: list[float], rng: random.Random) -> str:
    style = rng.random()
    if style < 0.5:
        return " x ".join(f"{d:g}" for d in dims) + " cm"
    if style < 0.75:
        return "x".join(f"{d * 10:g}" for d in dims) + "mm"
    return " x ".join(f"{d / 2.54:.1f}" for d in dims) + " in"


def _weight_a(grams: int, rng: random.Random) -> str:
    style = rng.random()
    if style < 0.4:
        return f"{grams} g"
    if style < 0.75:
        return f"{grams / 1000:.2f}kg".replace(".", ",")
    return f"{grams / 453.592:.2f} lb"


_ABBR = {"white": "WHT", "black": "BLK", "navy": "NVY", "gray": "GRY", "red": "RED", "green": "GRN", "olive": "OLV"}


def _title_a(p: GoldProduct, size: str | None, rng: random.Random) -> str:
    parts = [p.title]
    if p.feed_pattern:
        parts.insert(0, p.feed_pattern.title())
    if p.feed_count:
        parts.append(f"{p.feed_count}-Pack" if p.kind in {"socks", "towels", "candles"} else f"Set of {p.feed_count}")
    if p.feed_color:
        parts.append(_ABBR.get(p.feed_color, p.feed_color.upper()) if rng.random() < 0.5 else p.feed_color.title())
    if size:
        parts.append(size)
    title = " ".join(parts)
    return title.upper() if rng.random() < 0.4 else title


def _spec_notes(p: GoldProduct) -> str:
    bits = [*p.features]
    if p.origin:
        bits.append(f"Made in {p.origin}")
    return "; ".join(bits)


def write_supplier_a(products: list[GoldProduct], path: Path, rng: random.Random) -> list[dict[str, Any]]:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Price list"
    ws.append(["Larkspur Wholesale GmbH - autumn price list (net, EUR)"])
    ws.append(["Prices valid until 31.12.2026. All dimensions approx."])
    ws.append([])
    header = ["Item Name", "SKU#", "EAN", "Colour", "Mat.", "Size", "Dims", "Wt", "Price (EUR)", "Qty", "Image", "Category", "Notes"]
    ws.append(header)
    rows: list[dict[str, Any]] = []
    for p in products:
        sizes = p.sizes or [""]
        for j, size in enumerate(sizes):
            material_cell = p.composition if p.feed_material else ""
            if "material_mismatch" in p.variant_problems and j == len(sizes) - 1:
                material_cell = "100% polyester" if p.material != "polyester" else "100% cotton"
            price = p.price_usd
            if "price_outlier" in p.variant_problems and j == 1:
                price = price * 10
            row = {
                "Item Name": _title_a(p, size or None, rng),
                "SKU#": p.skus[j] if rng.random() < 0.8 else p.skus[j].lower().replace("-", ""),
                "EAN": p.gtins[j],
                "Colour": (p.feed_color.title() if p.feed_color else "") if rng.random() < 0.85 or not p.feed_color else "",
                "Mat.": material_cell,
                "Size": size,
                "Dims": _dims_a(p.dims_cm, rng) if p.dims_cm else "",
                "Wt": _weight_a(p.weight_g, rng),
                "Price (EUR)": _eur(price, rng),
                "Qty": rng.choice((0, 4, 12, 25, 40, 60)),
                "Image": p.image,
                "Category": KINDS[p.kind].supplier_category,
                "Notes": ". ".join(
                    x.rstrip(".")
                    for x in [_spec_notes(p), *p.fluff, (f"Capacity {p.capacity_ml} ml" if p.capacity_ml else ""), p.care]
                    if x
                ),
            }
            rows.append(row)
    for row in rows:
        ws.append([row[h] for h in header])
    for col, width in zip("ABCDEFGHIJKLM", (38, 12, 15, 10, 26, 7, 20, 10, 12, 6, 14, 18, 70), strict=True):
        ws.column_dimensions[col].width = width
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return rows


def write_supplier_b(products: list[GoldProduct], path: Path, rng: random.Random) -> list[dict[str, Any]]:
    header = ["product_title", "variant_size", "colour_name", "material_composition", "length_cm", "width_cm",
              "height_cm", "weight_kg", "capacity", "rrp_usd", "stock", "picture", "features", "care", "gtin", "item_code"]
    rows: list[dict[str, Any]] = []
    for p in products:
        sizes = p.sizes or [""]
        for j, size in enumerate(sizes):
            title = p.title
            if p.feed_pattern:
                title = f"{p.title} ({p.feed_pattern})"
            if p.feed_count:
                title += f", pack of {p.feed_count}"
            material_cell = p.composition if p.feed_material else ""
            if "material_mismatch" in p.variant_problems and j == len(sizes) - 1:
                material_cell = "100% polyester" if p.material != "polyester" else "100% cotton"
            price = p.price_usd * (10 if "price_outlier" in p.variant_problems and j == 1 else 1)
            dims = p.dims_cm or []
            rows.append({
                "product_title": title.lower() if rng.random() < 0.3 else title,
                "variant_size": size,
                "colour_name": p.feed_color or "",
                "material_composition": material_cell,
                "length_cm": f"{dims[0]:g}" if len(dims) > 0 else "",
                "width_cm": "",
                "height_cm": f"{dims[1]:g}" if len(dims) > 1 else "",
                "weight_kg": f"{p.weight_g / 1000:.3f}",
                "capacity": f"{p.capacity_ml} ml" if p.capacity_ml else "",
                "rrp_usd": f"{price:.2f}",
                "stock": str(rng.choice((0, 5, 18, 30, 75))),
                "picture": f"photos/{p.image}",
                "features": " | ".join([*p.features, *([f"Made in {p.origin}"] if p.origin else []), *p.fluff]),
                "care": p.care,
                "gtin": p.gtins[j],
                "item_code": p.skus[j],
            })
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=header, delimiter=";")
        writer.writeheader()
        writer.writerows(rows)
    return rows


def seed_catalog(out_dir: Path, *, seed: int = 11, n: int = 64, images: bool = True) -> dict[str, Any]:
    """Write feeds, images and gold labels into `out_dir`; returns the gold document."""
    rng = random.Random(seed + 1)
    products = _plan(seed, n)
    a = [p for p in products if p.supplier == "A"]
    b = [p for p in products if p.supplier == "B"]
    # duplicates: three rows repeated within supplier A (with cosmetic differences), three supplier A products that
    # supplier B also lists (same GTIN, B's own SKU and title)
    dup_within = rng.sample([p for p in a if not p.sizes], k=3)
    dup_cross = rng.sample([p for p in a if not p.sizes and p not in dup_within], k=3)
    rows_a = write_supplier_a(a, out_dir / "feeds" / SUPPLIER_A, rng)
    rows_b = write_supplier_b(b, out_dir / "feeds" / SUPPLIER_B, rng)
    _append_duplicates(out_dir, rows_a, rows_b, dup_within, dup_cross, rng)
    if images:
        for p in products:
            save(
                RenderSpec(p.kind, p.color, p.secondary_color, p.pattern, p.material, p.count, seed=seed * 1000 + int(p.id[1:])),
                out_dir / "images" / p.image,
            )
    gold = {
        "seed": seed,
        "note": "Synthetic catalog generated by shelfwise.synth.catalog; every product, brand and EAN is fictional.",
        "suppliers": {"A": SUPPLIER_A, "B": SUPPLIER_B},
        "products": [asdict(p) for p in products],
        "duplicates": {
            "within_supplier_a": [p.id for p in dup_within],
            "across_suppliers": [p.id for p in dup_cross],
        },
    }
    (out_dir / "gold.json").write_text(json.dumps(gold, indent=1, ensure_ascii=False), encoding="utf-8")
    return gold


def _append_duplicates(out_dir: Path, rows_a: list[dict[str, Any]], rows_b: list[dict[str, Any]],
                       within: list[GoldProduct], cross: list[GoldProduct], rng: random.Random) -> None:
    from openpyxl import load_workbook

    path_a = out_dir / "feeds" / SUPPLIER_A
    wb = load_workbook(path_a)
    ws = wb.active
    assert ws is not None
    header = [c.value for c in ws[4]]
    for p in within:
        original = next(r for r in rows_a if r["EAN"] == p.gtins[0])
        copy = dict(original)
        copy["Item Name"] = "  " + str(original["Item Name"]).title() + " "
        copy["SKU#"] = str(original["SKU#"]).lower()
        copy["Qty"] = rng.choice((3, 7))
        ws.append([copy.get(str(h), "") for h in header])
    wb.save(path_a)
    path_b = out_dir / "feeds" / SUPPLIER_B
    with path_b.open("a", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter=";")
        for p in cross:
            dims = p.dims_cm or []
            writer.writerow([
                f"{p.brand} {p.title}", "", p.feed_color or p.color, p.composition or "",
                f"{dims[0]:g}" if dims else "", "", f"{dims[1]:g}" if len(dims) > 1 else "", f"{p.weight_g / 1000:.3f}",
                f"{p.capacity_ml} ml" if p.capacity_ml else "", f"{p.price_usd + 2:.2f}", "10", f"photos/{p.image}",
                " | ".join(p.features), p.care, p.gtins[0], f"NB-X{p.id[1:]}",
            ])
